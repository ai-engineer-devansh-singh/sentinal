"""Generate SENTINEL Architecture Diagram as a .docx for the team."""
from __future__ import annotations

import os
import tempfile
from pathlib import Path

import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

OUT_PATH = Path(__file__).resolve().parent / "SENTINEL_Architecture.docx"


def draw_architecture_diagram() -> str:
    """Draw the architecture diagram and return the PNG file path."""
    fig, ax = plt.subplots(figsize=(18, 14))
    ax.set_xlim(0, 18)
    ax.set_ylim(0, 14)
    ax.axis("off")
    ax.set_facecolor("#f8f9fa")
    fig.patch.set_facecolor("#f8f9fa")

    def draw_box(x, y, w, h, text, color, text_color="white", fontsize=10, bold=True):
        box = FancyBboxPatch(
            (x - w / 2, y - h / 2), w, h,
            boxstyle="round,pad=0.02,rounding_size=0.2",
            facecolor=color, edgecolor="none", linewidth=0, zorder=2
        )
        ax.add_patch(box)
        weight = "bold" if bold else "normal"
        ax.text(x, y, text, ha="center", va="center", fontsize=fontsize,
                color=text_color, fontweight=weight, zorder=3, wrap=True)
        return box

    def draw_arrow(x1, y1, x2, y2, color="#495057", style="-|>", lw=2):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle=style, color=color, lw=lw,
                                    connectionstyle="arc3,rad=0"),
                    zorder=1)

    # Title
    ax.text(9, 13.4, "SENTINEL Protocol — System Architecture", fontsize=20,
            fontweight="bold", ha="center", va="center", color="#212529")
    ax.text(9, 12.9, "Multi-Agent Sepsis Detection System  |  Erasmus Toulouse Demo  |  July 2026",
            fontsize=11, ha="center", va="center", color="#6c757d")

    # --- Layer 1: Device Simulator ---
    draw_box(9, 11.6, 7.0, 1.0, "Device Simulator (Wearable Patch + POCT i-STAT)",
             "#0d6efd", fontsize=11)
    draw_arrow(9, 11.1, 9, 10.4)

    # --- Layer 2: FastAPI Backend (big container) ---
    backend = FancyBboxPatch(
        (2.5, 2.8), 13.0, 7.2,
        boxstyle="round,pad=0.02,rounding_size=0.3",
        facecolor="#e9ecef", edgecolor="#adb5bd", linewidth=2, linestyle="--", zorder=0
    )
    ax.add_patch(backend)
    ax.text(9, 9.7, "FastAPI Backend (Python 3.10 + asyncio)", fontsize=13,
            fontweight="bold", ha="center", va="center", color="#495057")

    # Feature Store
    draw_box(3.5, 8.6, 3.2, 0.9, "Feature Store\nrolling 15-min window\nSOFA / qSOFA / EWMA", "#198754", fontsize=9)
    # Agents
    draw_box(7.2, 8.6, 2.4, 0.9, "SOFA Tracker\n(rule engine)", "#dc3545", fontsize=9)
    draw_box(10.0, 8.6, 2.4, 0.9, "Differential\n(rule engine)", "#dc3545", fontsize=9)
    draw_box(12.8, 8.6, 2.4, 0.9, "Predictor\n(XGBoost)", "#fd7e14", fontsize=9)
    draw_box(15.4, 8.6, 2.2, 0.9, "Narrative\n(TF-IDF + LogReg)", "#fd7e14", fontsize=9)
    # Arrows from feature store to agents
    draw_arrow(5.1, 8.6, 6.0, 8.6, lw=1.5)
    draw_arrow(5.1, 8.6, 8.8, 8.6, lw=1.5)
    draw_arrow(5.1, 8.6, 11.6, 8.6, lw=1.5)
    draw_arrow(5.1, 8.6, 14.3, 8.6, lw=1.5)

    # Carbon Budget Controller
    draw_box(9, 6.8, 4.0, 0.9, "Carbon Budget Controller\n(LLM tier selection by ambiguity)", "#6f42c1", fontsize=9)
    draw_arrow(7.2, 8.15, 8.0, 7.25, lw=1.5)
    draw_arrow(10.0, 8.15, 9.5, 7.25, lw=1.5)
    draw_arrow(12.8, 8.15, 11.0, 7.25, lw=1.5)
    draw_arrow(15.4, 8.15, 13.0, 7.25, lw=1.5)

    # Orchestrator
    draw_box(9, 5.0, 5.0, 1.1, "Orchestrator — 3-Stage Deliberation Council\nStage 1: Independent assessments  →  Stage 2: Cross-examination  →  Stage 3: Synthesis",
        "#0dcaf0", text_color="#000000", fontsize=9)
    draw_arrow(9, 6.35, 9, 5.55, lw=2)

    # Grok Client / Local Fallback
    draw_box(14.5, 5.0, 3.0, 1.1, "Grok Client\nxAI Grok-4 / Grok-4-fast\n+ Local Fallback", "#6610f2", fontsize=9)
    draw_arrow(11.5, 5.0, 13.0, 5.0, lw=1.5, style="<->")

    # Alert Payload
    draw_box(9, 3.5, 3.8, 0.9, "Alert Payload\nTier (GREEN/YELLOW/RED) + Consensus + Recommended Action",
             "#d63384", fontsize=9)
    draw_arrow(9, 4.45, 9, 3.95, lw=2)

    # --- Layer 3: Dashboard & Audit ---
    draw_box(4.5, 3.5, 2.8, 0.9, "WebSocket\nPub/Sub Broadcast", "#20c997", text_color="#000000", fontsize=9)
    draw_arrow(7.1, 3.5, 7.1, 3.5)  # stub
    draw_arrow(7.1, 3.5, 7.1, 3.5)

    # Dashboard
    draw_box(4.5, 1.8, 3.0, 1.0, "Dashboard\nVanilla HTML/CSS/JS\nOffline Chart.js", "#0d6efd", fontsize=10)
    draw_arrow(4.5, 3.05, 4.5, 2.3, lw=2)

    # Audit Trail
    draw_box(9, 1.8, 3.0, 1.0, "Audit Trail\nAppend-only JSONL\n(Governance / Checkpoint 4–5)", "#6c757d", fontsize=9)
    draw_arrow(9, 3.05, 9, 2.3, lw=2)

    # HITL Controls
    draw_box(13.5, 1.8, 3.0, 1.0, "HITL Controls\nAccept / Modify / Reject\nClinician Response", "#ffc107", text_color="#000000", fontsize=9)
    draw_arrow(13.5, 2.8, 13.5, 2.3, lw=2)

    # Arrow from Alert Payload to WebSocket
    draw_arrow(7.1, 3.5, 6.0, 3.5, lw=1.5)

    # Legend
    legend_x = 0.5
    legend_y = 4.0
    legend_items = [
        ("#0d6efd", "I/O & Frontend"),
        ("#198754", "Feature Engineering"),
        ("#dc3545", "Rule-Based Agent"),
        ("#fd7e14", "ML-Based Agent"),
        ("#6f42c1", "Controller"),
        ("#0dcaf0", "Orchestration"),
        ("#6610f2", "LLM / AI"),
        ("#d63384", "Alerting"),
        ("#6c757d", "Persistence"),
        ("#ffc107", "Human-in-the-Loop"),
    ]
    ax.text(legend_x, legend_y + 2.8, "Legend", fontsize=11, fontweight="bold", color="#212529", va="center")
    for i, (color, label) in enumerate(legend_items):
        y = legend_y + 2.4 - i * 0.45
        rect = mpatches.Rectangle((legend_x, y - 0.15), 0.35, 0.30, facecolor=color, edgecolor="none", zorder=5)
        ax.add_patch(rect)
        ax.text(legend_x + 0.5, y, label, fontsize=9, va="center", color="#343a40")

    # Data flow annotation
    ax.text(16.5, 11.0, "Data Flow", fontsize=11, fontweight="bold", ha="center", color="#212529")
    ax.text(16.5, 10.5, "Simulator → Feature Store → Agents → Orchestrator → Alert → WS → Dashboard", fontsize=8, ha="center", color="#6c757d", wrap=True)

    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    plt.tight_layout()
    plt.savefig(tmp.name, dpi=200, bbox_inches="tight", facecolor="#f8f9fa")
    plt.close(fig)
    return tmp.name


def create_docx(diagram_path: str) -> None:
    doc = Document()

    # Title
    title = doc.add_heading("SENTINEL Protocol — System Architecture", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.runs[0]
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x21, 0x52, 0x29)

    subtitle = doc.add_paragraph("Multi-Agent Sepsis Detection System  |  Team Architecture Document")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = RGBColor(0x6c, 0x75, 0x7d)

    doc.add_paragraph("Date: July 2026  |  Demo: Erasmus Toulouse  |  Status: Production-Ready Demo")
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].runs[0].font.size = Pt(10)
    doc.paragraphs[-1].runs[0].font.color.rgb = RGBColor(0x6c, 0x75, 0x7d)

    doc.add_paragraph()

    # Diagram
    doc.add_heading("Architecture Overview", level=1)
    p = doc.add_paragraph()
    p.add_run("The diagram below shows the full SENTINEL stack, from simulated device input through the multi-agent deliberation council to the live dashboard and audit trail.").font.size = Pt(10)
    doc.add_picture(diagram_path, width=Inches(6.5))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # Component breakdown
    doc.add_heading("Component Breakdown", level=1)

    components = [
        ("1. Device Simulator", "Simulates a wearable patch + POCT i-STAT device. Streams realistic vitals for 6 scripted patient test cases (Karen, Anya, Maria, Robert, James, Frank) with deterministic noise (seed=42). Reproducible demo every run."),
        ("2. Feature Store (feature_store.py)", "Rolling 15-minute window with slopes, deltas, EWMA smoothing, SOFA/qSOFA scoring (Sepsis-3 proxies), and shock index. Deque-based, zero-allocation updates."),
        ("3. Four Specialist Agents (agents.py)", """• SOFA Tracker — Sepsis-3 organ-dysfunction rule engine (deterministic).\n• Differential — Mimic rule-out for pancreatitis and post-op stress (deterministic).\n• Predictor — XGBoost classifier for P(sepsis), trained on 7,200 synthetic samples (AUC 1.00, recall 0.99).\n• Narrative Analyst — TF-IDF + LogisticRegression for P(note indicates concern)."""),
        ("4. Carbon Budget Controller (agents.py)", "Selects LLM tier by case ambiguity: <0.30 → ml_only (~0.01g CO₂), 0.30–0.70 → grok-4-fast (~0.15g CO₂), ≥0.70 → grok-4 (~1.20g CO₂)."),
        ("5. Orchestrator (orchestrator.py)", "3-stage deliberation council: Stage 1 (independent deterministic assessments), Stage 2 (cross-examination via Grok), Stage 3 (synthesis → weighted consensus + tier + recommended action). Hallucination-grounding: consensus and tier are always deterministic."),
        ("6. Grok Client (grok_client.py)", "Async OpenAI-compatible client pointing to xAI endpoint (https://api.x.ai/v1). Supports grok-4 and grok-4-fast. Deterministic local-reasoning fallback if no key / auth failure / timeout."),
        ("7. FastAPI Backend (main.py)", "Async-native ASGI server. REST endpoints for device control, HITL response, speed/pause, audit, and metrics. Native WebSocket (/ws) with per-subscriber asyncio.Queue for pub/sub broadcast."),
        ("8. Dashboard (static/)", "Self-contained vanilla HTML/CSS/JS. Offline charting (Chart.js vendored). Live vitals grid, per-patient drawer with agent reasoning cards, sparklines, SOFA breakdown, and HITL controls."),
        ("9. Audit Trail (audit.py)", "Append-only JSONL log. Every tier change, clinician response, and model invocation recorded. Supports governance requirements (Checkpoint 4–5)."),
        ("10. ML Training Pipeline (ml/)", "Synthetic cohort generator + XGBoost + TF-IDF trainer. Deterministic, ~10s runtime. Models persisted with joblib. Automatic rule fallback if model files are missing."),
    ]

    for heading, body in components:
        h = doc.add_heading(heading, level=2)
        h.runs[0].font.color.rgb = RGBColor(0x0d, 0x6e, 0xfd)
        p = doc.add_paragraph(body)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        for run in p.runs:
            run.font.size = Pt(10)
        doc.add_paragraph()

    doc.add_page_break()

    # Data Flow section
    doc.add_heading("Data Flow (End-to-End)", level=1)
    flow = doc.add_paragraph()
    flow.add_run("""Device Simulator → Feature Store → 4 Agents (Stage 1, zero-LLM) → Carbon Budget Controller → Orchestrator (Stage 2 cross-exam + Stage 3 synthesis via Grok) → Alert Payload → WebSocket Broadcast → Dashboard

Clinician Response (Accept / Modify / Reject) → Audit Trail + Outcome Tracker

The 80% zero-LLM path means most inferences never call an LLM — only ambiguous cases trigger Grok deliberation, keeping latency low and cost minimal.""")
    flow.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for run in flow.runs:
        run.font.size = Pt(10)

    doc.add_paragraph()

    # Tech Stack Table
    doc.add_heading("Tech Stack Summary", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Layer"
    hdr_cells[1].text = "Technology"
    hdr_cells[2].text = "Purpose"
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(10)

    rows = [
        ("Language", "Python 3.10", "Type hints, dataclasses, asyncio"),
        ("Backend", "FastAPI + Uvicorn", "REST + WebSocket, async-native ASGI"),
        ("LLM", "xAI Grok (grok-4 / grok-4-fast)", "3-stage deliberation + synthesis"),
        ("LLM SDK", "openai >=1.30", "OpenAI-compatible async client"),
        ("ML", "XGBoost + scikit-learn", "Risk classifier + note concern classifier"),
        ("Serialization", "joblib", "Model persistence + lazy load"),
        ("Validation", "Pydantic v2", "HTTP API request/response models"),
        ("Config", "python-dotenv", "Env-driven dataclass Settings"),
        ("Frontend", "Vanilla HTML/CSS/JS", "No build step, fully offline"),
        ("Charts", "Chart.js (vendored)", "Offline sparklines + vitals charts"),
        ("Persistence", "JSONL (append-only)", "Audit trail, no DB dependency"),
        ("Simulation", "Deterministic script + noise", "Reproducible patient cases"),
    ]

    for layer, tech, purpose in rows:
        row_cells = table.add_row().cells
        row_cells[0].text = layer
        row_cells[1].text = tech
        row_cells[2].text = purpose
        for cell in row_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)

    doc.add_paragraph()

    # Patient Test Cases
    doc.add_heading("Patient Test Cases", level=1)
    p = doc.add_paragraph("Six scripted cases stream simultaneously to demonstrate full system behaviour:")
    p.runs[0].font.size = Pt(10)

    cases = [
        ("Karen Whitfield", "ER", "Sepsis (headline)", "GREEN → YELLOW (~t=30) → RED (~t=44). Caught before collapse."),
        ("Anya Petrova", "ICU", "Septic shock", "RED early, de-escalates to YELLOW on recovery."),
        ("Maria Sanchez", "ER", "Pancreatitis mimic", "Stays GREEN — Differential rules out sepsis."),
        ("Robert Hayes", "Surgical", "Post-op stress mimic", "Stays GREEN — distinguished from sepsis."),
        ("James Okafor", "General", "Healthy control", "Stays GREEN — silent monitoring."),
        ("Frank Doyle", "Palliative", "Sepsis + Comfort Care directive", "Consensus ~88 (sepsis physiology) but tier capped at YELLOW. Goals-of-care pathway, not aggressive bundle. Ethics safeguard demo."),
    ]

    table2 = doc.add_table(rows=1, cols=4)
    table2.style = "Light Grid Accent 1"
    hdr = table2.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Patient", "Ward", "Case", "Expected Behaviour"
    for cell in hdr:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(10)

    for patient, ward, case, behaviour in cases:
        row_cells = table2.add_row().cells
        row_cells[0].text = patient
        row_cells[1].text = ward
        row_cells[2].text = case
        row_cells[3].text = behaviour
        for cell in row_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)

    doc.add_paragraph()

    # Key Design Principles
    doc.add_heading("Key Design Principles", level=1)
    principles = [
        "Deterministic Safety Layer — Consensus and alert tier are always rule-derived; the LLM only explains. Alerts fire correctly even if Grok is offline.",
        "Carbon-Aware AI — The Carbon Budget Controller selects model size by ambiguity. 80% of cases take the zero-LLM path (~0.01g CO₂).",
        "Graceful Degradation — Local-reasoning fallback stubs produce clinician-style explanations when xAI is unavailable. Demo never breaks.",
        "Human-in-the-Loop — Every RED/YELLOW alert requires clinician response (Accept/Modify/Reject) with counterfactual outcome tracking.",
        "Ethics Safeguard — Advance directives (e.g., Comfort Care) cap alert tier and suppress aggressive bundles, preserving patient dignity.",
        "Reproducibility — Deterministic simulator with fixed noise seed means every demo run is identical and predictable.",
    ]
    for principle in principles:
        p = doc.add_paragraph(principle, style="List Bullet")
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        for run in p.runs:
            run.font.size = Pt(10)

    doc.add_paragraph()
    doc.add_paragraph("— Generated by SENTINEL build tooling  |  For internal team use only")
    doc.paragraphs[-1].runs[0].font.size = Pt(9)
    doc.paragraphs[-1].runs[0].font.color.rgb = RGBColor(0x6c, 0x75, 0x7d)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.save(str(OUT_PATH))
    print(f"Saved architecture document to: {OUT_PATH}")


if __name__ == "__main__":
    diagram_path = draw_architecture_diagram()
    try:
        create_docx(diagram_path)
    finally:
        os.unlink(diagram_path)
